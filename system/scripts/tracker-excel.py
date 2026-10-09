# Build "Social Accords Tracker.xlsx" from the posting tracker's exported rows.
# Input:  system/projects/_metricool-upload/tracker/export/videos/*.json  (ArtifactData list with out_dir)
# Output: system/projects/_metricool-upload/tracker/Social Accords Tracker.xlsx
import json, glob, os, re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = "C:/Users/taldo/Downloads/videos to edit/system/projects/_metricool-upload/tracker"
V = sorted([json.load(open(f, encoding="utf8")) for f in glob.glob(ROOT + "/export/videos/*.json")], key=lambda d: d["n"])
# Instagram reel code -> (views, shares, saves), read from Metricool on 2026-10-10
A = {"DeNVi0coo77": (31776, 176, 108), "DeMlEjijOav": (118285, 2653, 1169), "DeMGp-hjji2": (154153, 2625, 1181),
     "DeKdfwTIlWS": (205959, 1256, 917), "DeITK8iIrub": (386981, 3458, 2471), "DeDBrNIosDd": (372303, 20474, 2551),
     "DeAYTFFtubz": (28300, 29, 60), "Dd9fyZDICq0": (149687, 360, 179), "Dd60aXOo8Ku": (479905, 3672, 1258),
     "Ddw_hauoMDq": (913740, 5889, 3748), "Ddr6p4WIK_n": (391707, 2208, 1246), "DdeLJhLIEN5": (51397, 107, 79),
     "DdU9f5Po5Jv": (43908, 210, 124), "DdR_ofzI7Il": (37762, 91, 98)}
HF = Font(bold=True, color="FFFFFF"); HFILL = PatternFill("solid", fgColor="14212B")
OK = PatternFill("solid", fgColor="E0F2E4"); SC = PatternFill("solid", fgColor="FBEFD2")
NO = PatternFill("solid", fgColor="F1F4F6"); WARN = PatternFill("solid", fgColor="FBE6DF")
B = Border(bottom=Side(style="thin", color="D6DEE4"))
LINK = Font(color="0B6E75", underline="single")
PL = (("ig", "Instagram"), ("tt", "TikTok"), ("yt", "YouTube"), ("x", "X"))
MON = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def head(ws, cols, widths, freeze="A2"):
    ws.append(cols)
    for i, w in enumerate(widths, 1):
        x = ws.cell(1, i); x.font = HF; x.fill = HFILL; x.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30; ws.freeze_panes = freeze


def finish(ws):
    ws.auto_filter.ref = ws.dimensions
    for r in ws.iter_rows(min_row=2):
        for c in r:
            c.border = B; c.alignment = Alignment(vertical="top", wrap_text=True)


def pick(ws, rng, options):
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    ws.add_data_validation(dv); dv.add(rng)


def fmt(d):
    m = re.match(r"(\d{4})-(\d\d)-(\d\d) ?(.*)", d or "")
    return f"{int(m.group(3))} {MON[int(m.group(2))]} {m.group(4)}".strip() if m else (d or "")


def stage(v):
    if v["edit"] != "Ready to Post": return v["edit"]
    p = sum(v[k]["s"] == "posted" for k, _ in PL)
    return "Posted everywhere" if p == 4 else ("Partly posted" if p else "Ready, not posted")


wb = Workbook()
ds = wb.active; ds.title = "Dashboard"

# ---- Videos
ws = wb.create_sheet("Videos")
head(ws, ["#", "Video", "Frame.io file", "Editing status", "Location", "Instagram link", "Instagram views", "Shares", "Saves",
          "Instagram", "TikTok", "YouTube", "X", "Posted where", "Frame.io link", "Notes"],
     [5, 44, 40, 16, 11, 48, 12, 9, 9, 16, 18, 18, 18, 20, 14, 46], "C2")
for v in V:
    m = re.search(r"/reel/([^/]+)/", v["ig"].get("url", ""))
    a = A.get(m.group(1)) if m else None
    cells = ["✅ Posted" if v[k]["s"] == "posted" else ("🗓️ " + fmt(v[k].get("at", "")) if v[k]["s"] == "scheduled" else "⬜ Not scheduled") for k, _ in PL]
    missing = v["edit"] == "Ready to Post" and stage(v) != "Posted everywhere" and not v.get("frameName")
    ws.append([v["n"], v["title"], v.get("frameName") or ("Missing from Frame.io" if missing else ""), v["edit"], v.get("loc", ""),
               v["ig"].get("url", ""), a[0] if a else (v["views"] or None), a[1] if a else None, a[2] if a else None,
               *cells, stage(v), "Open" if v.get("frame") else "", v.get("note", "")])
    r = ws.max_row
    if v["ig"].get("url"): ws.cell(r, 6).hyperlink = v["ig"]["url"]; ws.cell(r, 6).font = LINK
    if v.get("frame"): ws.cell(r, 15).hyperlink = v["frame"]; ws.cell(r, 15).font = LINK
    for i, (k, _) in enumerate(PL):
        c = ws.cell(r, 10 + i); c.fill = {"posted": OK, "scheduled": SC}.get(v[k]["s"], NO)
        if v[k].get("url"): c.hyperlink = v[k]["url"]
    if missing: ws.cell(r, 3).fill = WARN
    for c in (7, 8, 9): ws.cell(r, c).number_format = "#,##0"
finish(ws)
pick(ws, "D2:D200", ["Not Filmed", "Not Edited", "In Editing", "Ready to Post"])
n = ws.max_row

# ---- Calendar
cal = wb.create_sheet("Calendar")
head(cal, ["Date", "Time", "Platform", "Video", "Status", "Link once live"], [14, 9, 22, 48, 16, 48])
rows = []
for v in V:
    for k, l in PL:
        p = v[k]
        if p["s"] == "scheduled" and p.get("at"):
            d, _, t = p["at"].partition(" "); rows.append((d, t, l, v["title"]))
rows += [("2026-10-10", "16:30", "Instagram trial reel", "Shop Owner 5: Juice Kindness Test"),
         ("2026-10-15", "16:00", "Instagram trial reel", "Flowers 2: A Gift to Make Your Day Better")]
for d, t, l, ti in sorted(rows):
    cal.append([d, t, l, ti, "Scheduled", ""]); cal.cell(cal.max_row, 5).fill = SC
finish(cal)
pick(cal, "E2:E500", ["Scheduled", "Posted", "Cancelled"])

# ---- Creator collabs
cc = wb.create_sheet("Creator Collabs")
head(cc, ["Creator", "Instagram handle", "Video", "Status", "Instagram link", "Views", "Location", "Next step / notes"], [26, 22, 44, 14, 48, 12, 12, 46])
byn = {v["n"]: v for v in V}
for name, h, nn in [("The Traveling Clatt", "@thetravelingclatt", 1), ("Lielle Blinkoff", "@lielleblinkoff", 2), ("Gabe Einhorn", "", 3),
                    ("Eliya Cohen", "", 4), ("Gabe Einhorn", "", 7), ("Niall Donnan", "@niall.donnan", 14), ("Montana Tucker", "@montanatucker", 15),
                    ("Montana Tucker", "@montanatucker", 16), ("Kowshee", "@__kowshee__", 20), ("Ben Gon", "@ben_gon", 25),
                    ("Montana Tucker", "@montanatucker", 29), ("Zaddy Yellow", "@zaddy_yellow_ent", 30)]:
    v = byn[nn]; need = [l for k, l in PL[1:] if v[k]["s"] == "none"]
    note = ("Still needs " + ", ".join(need) + (". File missing from Frame.io." if not v.get("frame") else ".")) if need else "Posted or scheduled everywhere."
    cc.append([name, h, v["title"], "Done", v["ig"].get("url", ""), v["views"] or None, v.get("loc", ""), note])
    r = cc.max_row; cc.cell(r, 4).fill = OK; cc.cell(r, 6).number_format = "#,##0"
    if v["ig"].get("url"): cc.cell(r, 5).hyperlink = v["ig"]["url"]; cc.cell(r, 5).font = LINK
for row in [("Montana Tucker + James Maslow", "@montanatucker", "Phone-call video", "In editing", "", None, "", "Still editing."),
            ("Montana Tucker", "@montanatucker", "Dance video in LA", "In editing", "", None, "LA", "Filmed, almost finished. Waiting on the final edit."),
            ("Yoav Beatbox", "", "Having a Bad Day / You're Sad (sit-down)", "Filmed", "", None, "Israel", "Filmed 10 October. Needs editing.")]:
    cc.append(list(row)); cc.cell(cc.max_row, 4).fill = SC
for _ in range(12): cc.append(["", "", "", "Idea", "", None, "", ""])
finish(cc)
pick(cc, "D2:D300", ["Idea", "Reached out", "Confirmed", "Filmed", "In editing", "Done"])

# ---- Shoots
sh = wb.create_sheet("Shoots")
head(sh, ["Shoot", "Location", "Date", "Status", "Who is in it", "Format", "Editor", "Frame.io folder", "Notes"], [42, 12, 14, 16, 26, 24, 16, 38, 46])
for row in [("Montana's dance video", "LA", "", "In Editing", "Montana Tucker", "Dance", "", "", "Almost finished. Waiting on the final edit."),
            ("Montana with James Maslow: phone call", "", "", "In Editing", "Montana Tucker, James Maslow", "Call someone you love", "", "", ""),
            ("Yoav Beatbox: Having a Bad Day / You're Sad", "Israel", "2026-10-10", "Not Edited", "Tal, Yoav Beatbox", "Sit-down", "", "", ""),
            ("Free Hugs sign", "Israel", "", "Not Edited", "Tal", "Street kindness", "", "NEEDS TO BE EDITED / HUG FOR STRANGER", ""),
            ("Call Someone You Love interview", "Israel", "", "Not Edited", "Tal", "Call someone you love", "", "NEEDS TO BE EDITED / PHONE CALL VIDEO", "")]:
    sh.append(list(row))
for _ in range(15): sh.append(["", "", "", "Planned", "", "", "", "", ""])
finish(sh)
pick(sh, "D2:D300", ["Idea", "Planned", "Filmed", "Not Edited", "In Editing", "Ready to Post", "Posted"])
pick(sh, "B2:B300", ["Israel", "LA", "Miami", "New York", "Other"])

# ---- Ideas
idw = wb.create_sheet("Ideas")
head(idw, ["Idea", "Format", "Why (from your numbers)", "Location", "Who", "Priority", "Status", "Notes"], [48, 26, 60, 12, 22, 10, 14, 36])
for i in [
    ("POV: Meeting a Druze / Bedouin / Ethiopian Jew / Armenian in Israel", "POV meeting", "Arab Christian (914K) and Police Officer (387K) are this format: one identity in the title, a warm face on the cover.", "Israel", "Tal", "High"),
    ("An unlikely pair who love each other: three more couples or best friends", "Story, 60 to 100 seconds", "Fabian (Israeli and Lebanese couple) was shared 20,474 times, 5.5% of viewers. No other video is close.", "Israel", "Tal", "High"),
    ("Heroes: someone from one group who saved someone from another", "Story with interview", "Hisham and Aya, and Younes: 2.2% and 1.7% of viewers shared them, far above the rest.", "Israel", "Tal", "High"),
    ("I had no money: bakery, taxi, barber, falafel, pharmacy", "Kindness test with identity reveal", "Free Food 480K, Coffee Shop 392K. The reveal and 'we are brothers' is the moment people share.", "Israel", "Tal", "High"),
    ("Montana joins a POV: Meeting a ___ in LA or Miami", "POV meeting with Montana", "Her two Miami posts did 555K and 559K on her page. Bring your winning format to her audience.", "LA / Miami", "Montana, Tal", "High"),
    ("One creator from each Abraham Accords country", "Creator collab", "The Kowshee collab did 2.6M. Collabs on big pages beat everything on your own page.", "Remote", "Tal", "High"),
    ("So Much in Common, part 2: Muslim and Jewish strangers compare childhoods", "Two-person conversation", "Gabe's similarities video: 206K views, 30% average watch.", "Israel", "Tal, Gabe", "Medium"),
    ("The pledge: one video asking people to sign and tag a friend", "Call to action", "The pledge form and the DM automation exist. No video sends people to them yet.", "Any", "Montana, Tal", "Medium"),
    ("A blessing from elders: survivors, a sheikh, a priest", "Short portrait", "Sarah's blessing did 146K.", "Israel", "Tal", "Medium"),
    ("Kindness test with no identity in the title", "Street experiment", "Lower priority. Market test 28K and Who Stops to Help 150K had 0.1 to 0.2% shares.", "Israel", "Tal", "Low"),
]:
    idw.append(list(i) + ["Idea", ""])
for _ in range(20): idw.append(["", "", "", "", "", "", "Idea", ""])
finish(idw)
pick(idw, "F2:F300", ["High", "Medium", "Low"]); pick(idw, "G2:G300", ["Idea", "Planned", "Filmed", "Dropped"])

# ---- Files needed
mf = wb.create_sheet("Files Needed")
head(mf, ["#", "Video", "Instagram link", "Posted on", "Still needs", "Received?"], [5, 46, 50, 22, 26, 12])
for v in V:
    if v["edit"] == "Ready to Post" and not v.get("frame") and any(v[k]["s"] == "none" for k, _ in PL[1:]):
        u = v["ig"].get("url", ""); m = re.search(r"instagram.com/([^/]+)/", u)
        mf.append([v["n"], v["title"], u, "@" + m.group(1) if m else "", ", ".join(l for k, l in PL[1:] if v[k]["s"] == "none"), "No"])
        if u: mf.cell(mf.max_row, 3).hyperlink = u; mf.cell(mf.max_row, 3).font = LINK
finish(mf)
pick(mf, "F2:F100", ["No", "Yes"])

# ---- Dashboard (formulas, so the counts follow the Videos sheet)
e = n + 60
ds["A1"] = "Social Accords video tracker"; ds["A1"].font = Font(bold=True, size=18)
ds["A2"] = "Counts follow the Videos sheet. Data as of 10 October 2026."
rows = [("Total videos", f"=COUNTA(Videos!B2:B{e})"), ("Not filmed", f'=COUNTIF(Videos!D2:D{e},"Not Filmed")'),
        ("Filmed, needs editing", f'=COUNTIF(Videos!D2:D{e},"Not Edited")'), ("Being edited", f'=COUNTIF(Videos!D2:D{e},"In Editing")'),
        ("Ready, not posted", f'=COUNTIF(Videos!N2:N{e},"Ready, not posted")'), ("Partly posted", f'=COUNTIF(Videos!N2:N{e},"Partly posted")'),
        ("Posted everywhere", f'=COUNTIF(Videos!N2:N{e},"Posted everywhere")'), ("", ""),
        ("On Instagram", f'=COUNTIF(Videos!J2:J{e},"*Posted")'), ("On TikTok", f'=COUNTIF(Videos!K2:K{e},"*Posted")'),
        ("On YouTube", f'=COUNTIF(Videos!L2:L{e},"*Posted")'), ("On X", f'=COUNTIF(Videos!M2:M{e},"*Posted")'),
        ("Scheduled posts", '=COUNTIF(Calendar!E2:E500,"Scheduled")'), ("Files still needed", "=COUNTIF('Files Needed'!F2:F100,\"No\")"), ("", ""),
        ("20-video commitment: posted everywhere", '=B10&" of 20"')]
for i, (a, b) in enumerate(rows, 4):
    ds.cell(i, 1, a); ds.cell(i, 2, b); ds.cell(i, 1).font = Font(bold=bool(a)); ds.cell(i, 2).alignment = Alignment(horizontal="right")
ds.column_dimensions["A"].width = 44; ds.column_dimensions["B"].width = 14
ds["A22"] = "Sheets"; ds["A22"].font = Font(bold=True)
for i, t in enumerate(["Videos: one row per video. Green = posted, yellow = scheduled with the date, grey = not scheduled.",
                       "Calendar: every scheduled post in date order.", "Creator Collabs: done and upcoming.",
                       "Shoots: what is filmed and what is planned, by location.", "Ideas: what to film next, with the reason from your numbers.",
                       "Files Needed: what to upload to Frame.io."], 23):
    ds.cell(i, 1, t)
out = ROOT + "/Social Accords Tracker.xlsx"
wb.save(out)
print(out, os.path.getsize(out), "videos:", n - 1, "files needed:", mf.max_row - 1, "calendar rows:", cal.max_row - 1)
